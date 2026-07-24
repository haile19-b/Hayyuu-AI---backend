import logging
from io import BytesIO
import docx
from pypdf import PdfReader
from google import genai
from google.genai import types
import anyio

from app.core.env import settings
from app.core.config import genAI

logger = logging.getLogger("uvicorn.error")

# Configure accurate table parsing options for IBM Docling
try:
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling.datamodel.base_models import InputFormat, DocumentStream
    from docling.datamodel.pipeline_options import PdfPipelineOptions, TableFormerMode
    
    pipeline_options = PdfPipelineOptions(do_table_structure=True)
    pipeline_options.table_structure_options.mode = TableFormerMode.ACCURATE
    doc_converter = DocumentConverter(
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options)
        }
    )
    logger.info("✅ IBM Docling DocumentConverter Initialized Successfully")
except Exception as doc_err:
    logger.error(f"Failed to initialize Docling converter: {doc_err}")
    doc_converter = None


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Synchronous helper to extract text from a DOCX file."""
    try:
        doc = docx.Document(BytesIO(file_bytes))
        text = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text)
        return "\n".join(text)
    except Exception as e:
        logger.error(f"Error reading DOCX file: {e}")
        raise e


def extract_text_from_pdf_digital(file_bytes: bytes) -> str:
    """Synchronous helper to extract text from a digital PDF."""
    try:
        reader = PdfReader(BytesIO(file_bytes))
        text = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text.append(extracted)
        return "\n".join(text)
    except Exception as e:
        logger.error(f"Error reading PDF digitally: {e}")
        raise e


async def extract_text_via_gemini(file_bytes: bytes, mime_type: str) -> str:
    """Call Gemini to extract text from scanned PDFs or images using the google-genai SDK."""
    try:
        prompt = (
            "You are a document transcription system. Extract and transcribe all the textual content "
            "from this document. Present it clearly, preserving structural flow where possible. "
            "Do not summarize or add commentary. Output only the extracted text."
        )
        
        # Wrap bytes in the SDK Part model
        binary_part = types.Part.from_bytes(
            data=file_bytes,
            mime_type=mime_type
        )
        
        def _generate():
            response = genAI.models.generate_content(
                model="gemini-3.5-flash",
                contents=[binary_part, prompt]
            )
            return response.text
            
        return await anyio.to_thread.run_sync(_generate)
    except Exception as e:
        logger.error(f"Error performing OCR text extraction via Gemini: {e}")
        raise e


async def extract_text(file_bytes: bytes, content_type: str) -> str:
    """
    Core entry point to extract text based on mime-type.
    Uses IBM Docling for layout-aware Markdown extraction (tables, headers) when available.
    Falls back to legacy extractors on failure or unsupported formats.
    """
    c_type = content_type.lower()
    is_pdf = "pdf" in c_type or c_type == "application/pdf"
    is_docx = (
        "wordprocessingml.document" in c_type
        or c_type == "docx"
        or "vnd.openxmlformats-officedocument.wordprocessingml" in c_type
    )

    # 1. Attempt structured extraction via IBM Docling
    if doc_converter and (is_pdf or is_docx):
        logger.info(f"Attempting layout-aware text extraction for {content_type} via Docling")
        try:
            def _convert():
                ext = "pdf" if "pdf" in content_type.lower() else "docx"
                source = DocumentStream(name=f"document.{ext}", stream=BytesIO(file_bytes))
                result = doc_converter.convert(source)
                return result.document.export_to_markdown()
            
            markdown_text = await anyio.to_thread.run_sync(_convert)
            if markdown_text.strip():
                logger.info("Successfully extracted text via Docling.")
                return markdown_text
        except Exception as e:
            logger.warning(f"Docling extraction failed: {e}. Falling back to legacy extractors.")

    # 2. Legacy/Fallback extraction paths
    # DOCX Fallback
    if is_docx:
        logger.info("Extracting text from DOCX document using fallback helper")
        return await anyio.to_thread.run_sync(extract_text_from_docx, file_bytes)
        
    # PDF Fallback
    elif is_pdf:
        logger.info("Extracting text from PDF document using fallback helper")
        # Try digital extraction first
        digital_text = await anyio.to_thread.run_sync(extract_text_from_pdf_digital, file_bytes)
        
        # If we got reasonable text, return it
        if len(digital_text.strip()) > 100:
            logger.info("Successfully extracted text digitally from PDF")
            return digital_text
        else:
            logger.info("Digital PDF extraction returned minimal text. Falling back to Gemini OCR")
            return await extract_text_via_gemini(file_bytes, "application/pdf")
            
    # Images
    elif "image" in c_type or c_type in ["png", "jpg", "jpeg", "webp"]:
        logger.info(f"Extracting text from image ({content_type}) via Gemini Vision")
        mime_type = content_type if "image" in c_type else f"image/{c_type.replace('jpg', 'jpeg')}"
        return await extract_text_via_gemini(file_bytes, mime_type)
        
    # Text fallback
    else:
        logger.info(f"Treating content type {content_type} as plain text")
        try:
            return file_bytes.decode("utf-8")
        except Exception as e:
            logger.error(f"Failed to decode text file: {e}")
            raise ValueError(f"Unsupported content type for text extraction: {content_type}")
